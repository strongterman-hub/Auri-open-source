from __future__ import annotations

import io


# Keep injected file text bounded so a huge document cannot blow up the prompt.
MAX_EXTRACTED_CHARS = 40_000

TEXT_MIMES = {
    "application/json",
    "application/xml",
    "application/x-yaml",
    "text/csv",
    "text/tab-separated-values",
}

TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".markdown",
    ".csv",
    ".tsv",
    ".log",
    ".json",
    ".xml",
    ".yaml",
    ".yml",
    ".py",
    ".js",
    ".ts",
    ".java",
    ".kt",
    ".sh",
    ".ini",
    ".cfg",
    ".conf",
    ".html",
    ".css",
}


def _decode_text(data: bytes) -> str | None:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return None


def _extract_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _extract_xlsx(data: bytes) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows(values_only=True):
            cells = ["" if cell is None else str(cell) for cell in row]
            if any(cell.strip() for cell in cells):
                lines.append("\t".join(cells))
    return "\n".join(lines)


def _truncate(text: str) -> str:
    text = text.strip()
    if len(text) <= MAX_EXTRACTED_CHARS:
        return text
    return text[:MAX_EXTRACTED_CHARS] + "\n…[文件内容过长，已截断]"


def extract_text(name: str, mime: str, data: bytes) -> str | None:
    """Extract readable text from common document types, or return None."""
    mime = (mime or "").lower()
    name = (name or "").lower()

    if mime.startswith("text/") or mime in TEXT_MIMES:
        text = _decode_text(data)
        return _truncate(text) if text is not None else None

    if mime == "application/pdf" or name.endswith(".pdf"):
        try:
            return _truncate(_extract_pdf(data))
        except Exception:
            return None

    if (
        mime
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or name.endswith(".docx")
    ):
        try:
            return _truncate(_extract_docx(data))
        except Exception:
            return None

    if (
        mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        or name.endswith(".xlsx")
    ):
        try:
            return _truncate(_extract_xlsx(data))
        except Exception:
            return None

    # Last resort for plain-text files that arrived with a generic mime type.
    extension = name[name.rfind(".") :] if "." in name else ""
    if extension in TEXT_EXTENSIONS:
        text = _decode_text(data)
        return _truncate(text) if text is not None else None

    return None
